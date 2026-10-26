from io import BytesIO
from pathlib import Path

def extract_text(filename: str, data: bytes):
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            text = "\n".join(page.get_text("text") for page in doc).strip()
        except Exception as exc:
            raise ValueError(f"PDF parsing failed: {exc}")
    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts).strip()
        except Exception as exc:
            raise ValueError(f"DOCX parsing failed: {exc}")
    elif ext == ".txt":
        text = data.decode("utf-8", errors="replace").strip()
    else:
        raise ValueError("Unsupported format. Use PDF, DOCX, or TXT.")
    if len(text) < 30:
        return text, {"status": "Poor / limited extraction", "message": "Very little text was extracted. Check whether the PDF is scanned or the file is empty."}
    return text, {"status": "Text extracted", "message": f"{len(text):,} characters extracted."}
